from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ModuleNotFoundError as exc:
    raise SystemExit(
        "Не встановлено бібліотеку python-docx.\n"
        "Виконайте команду: py -3.13 -m pip install python-docx\n"
        "Якщо у системі доступний тільки python, використайте: python -m pip install python-docx"
    ) from exc


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output"
OUTPUT_FILE = OUTPUT_DIR / "PlantShop_qualification_work_full.docx"

FONT_NAME = "Times New Roman"
CODE_FONT = "Courier New"


def sanitize_text(text: str) -> str:
    return text.replace(chr(0x2014), "–")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size_pt: float = 14, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_code_run_font(run) -> None:
    run.font.name = CODE_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size in [
        ("Heading 1", 14),
        ("Heading 2", 14),
        ("Heading 3", 14),
    ]:
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    if "Code" not in [style.name for style in styles]:
        styles.add_style("Code", 1)
    code_style = styles["Code"]
    code_style.font.name = CODE_FONT
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
    code_style.font.size = Pt(12)
    code_style.paragraph_format.first_line_indent = Cm(0)
    code_style.paragraph_format.line_spacing = 1
    code_style.paragraph_format.space_after = Pt(0)

    add_page_number(section.header.paragraphs[0])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, 12)


def add_plain_paragraph(
    doc: Document,
    text: str,
    *,
    first_line: bool = True,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_after: int = 0,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.25 if first_line else 0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(sanitize_text(text))
    set_run_font(run, 14)


def add_centered(doc: Document, text: str, *, size: float = 14, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(sanitize_text(text))
    set_run_font(run, size, bold)


def add_right(doc: Document, text: str, *, size: float = 14, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(sanitize_text(text))
    set_run_font(run, size, bold)


def add_chapter_heading(doc: Document, chapter: str, title: str) -> None:
    doc.add_page_break()
    p1 = doc.add_paragraph()
    p1.style = doc.styles["Heading 1"]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.first_line_indent = Cm(0)
    p1.paragraph_format.space_after = Pt(6)
    run1 = p1.add_run(chapter.upper())
    set_run_font(run1, 14, True)

    p2 = doc.add_paragraph()
    p2.style = doc.styles["Heading 1"]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Cm(0)
    p2.paragraph_format.space_after = Pt(12)
    run2 = p2.add_run(sanitize_text(title.upper()))
    set_run_font(run2, 14, True)


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(sanitize_text(text))
    set_run_font(run, 14, True)


def add_block_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 3"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(sanitize_text(text))
    set_run_font(run, 14, True)


def add_figure_placeholder(doc: Document, number: str, title: str) -> None:
    add_plain_paragraph(
        doc,
        f"Зміст відповідної схеми або екрана наведено на рисунку {number}.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"[МІСЦЕ ДЛЯ РИСУНКА {number} – {title}]")
    set_run_font(run, 14, True)

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1.5
    caption.paragraph_format.space_after = Pt(8)
    run_caption = caption.add_run(f"Рисунок {number} – {title}")
    set_run_font(run_caption, 14)


def add_table(
    doc: Document,
    caption: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[float] | None = None,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(sanitize_text(caption))
    set_run_font(run, 14)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = header_cells[i]
        set_cell_shading(cell, "EDEDED")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            set_cell_width(cell, widths[i])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run_header = paragraph.add_run(sanitize_text(header))
        set_run_font(run_header, 11, True)
    set_repeat_table_header(table.rows[0])

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                set_cell_width(cell, widths[i])
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(0)
            run_value = paragraph.add_run(sanitize_text(str(value)))
            set_run_font(run_value, 10.5)

    doc.add_paragraph()


def add_listing(doc: Document, number: str, title: str, code: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    for line in code.strip("\n").splitlines():
        run = p.add_run(line.rstrip())
        set_code_run_font(run)
        run.add_break()

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1.5
    caption.paragraph_format.space_after = Pt(6)
    run_caption = caption.add_run(f"Лістинг {number} – {title}")
    set_run_font(run_caption, 14)


def add_toc_field(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "separate")
    fld_char_3 = OxmlElement("w:fldChar")
    fld_char_3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    text_run = p.add_run("Зміст буде оновлено засобами Microsoft Word після команди «Оновити поле».")
    set_run_font(text_run, 14)
    run._r.append(fld_char_3)


def add_title_page(doc: Document) -> None:
    add_centered(doc, "[НАЗВА ЗАКЛАДУ ВИЩОЇ ОСВІТИ]", bold=True)
    add_centered(doc, "[ФАКУЛЬТЕТ / ІНСТИТУТ]", bold=True)
    add_centered(doc, "[КАФЕДРА]", bold=True)
    for _ in range(6):
        doc.add_paragraph()
    add_centered(doc, "КВАЛІФІКАЦІЙНА РОБОТА БАКАЛАВРА", bold=True)
    add_centered(
        doc,
        "на тему:",
    )
    add_centered(
        doc,
        "«Розробка веб-додатку PlantShop для продажу рослин та надання послуг ландшафтного дизайну із використанням сучасних технологій Next.js та NestJS»",
        bold=True,
    )
    for _ in range(4):
        doc.add_paragraph()
    add_right(doc, "Виконав: студент групи [ГРУПА]")
    add_right(doc, "[ПРІЗВИЩЕ ІМ’Я ПО БАТЬКОВІ]")
    add_right(doc, "Керівник: [ПІБ, науковий ступінь, посада]")
    for _ in range(8):
        doc.add_paragraph()
    add_centered(doc, "[МІСТО] – 2026")
    doc.add_page_break()


def add_assignment(doc: Document) -> None:
    add_section_heading(doc, "ЗАВДАННЯ НА КВАЛІФІКАЦІЙНУ РОБОТУ БАКАЛАВРА")
    add_plain_paragraph(
        doc,
        "Тема кваліфікаційної роботи бакалавра: «Розробка веб-додатку PlantShop для продажу рослин та надання послуг ландшафтного дизайну із використанням сучасних технологій Next.js та NestJS».",
    )
    add_plain_paragraph(
        doc,
        "Метою роботи є розробка та опис full-stack веб-додатку PlantShop, який забезпечує публічний каталог рослин, автентифікацію користувачів, оформлення замовлень авторизованими користувачами, особистий кабінет та адміністративну панель для керування основними даними системи.",
    )
    add_plain_paragraph(
        doc,
        "Вихідними даними для виконання роботи є реалізований репозиторій PlantShop, що складається з frontend-застосунку на Next.js та backend API на NestJS, а також Prisma-схеми бази даних MySQL, тестових файлів і конфігурацій запуску.",
    )
    add_table(
        doc,
        "Таблиця З.1 – Основні завдання кваліфікаційної роботи",
        ["№", "Завдання", "Очікуваний результат"],
        [
            ["1", "Проаналізувати предметну область онлайн-продажу рослин", "Сформульовано потреби користувачів та адміністратора"],
            ["2", "Визначити вимоги до PlantShop", "Описано ролі, функціональні та нефункціональні вимоги"],
            ["3", "Розробити frontend-частину", "Реалізовано storefront, каталог, кошик, dashboard та admin UI"],
            ["4", "Розробити backend API", "Реалізовано auth, plants, orders, media та site settings modules"],
            ["5", "Спроєктувати базу даних", "Описано Prisma-моделі та зв’язки між ними"],
            ["6", "Перевірити працездатність", "Підготовлено тестове покриття та план ручної перевірки"],
        ],
        widths=[1.0, 6.2, 8.0],
    )
    add_plain_paragraph(
        doc,
        "Остаточне оформлення титульних реквізитів, дат, підписів та формальних полів завдання потребує ручного уточнення відповідно до вимог конкретного закладу освіти.",
    )
    doc.add_page_break()


def add_annotation(doc: Document) -> None:
    add_section_heading(doc, "АНОТАЦІЯ")
    paragraphs = [
        "Кваліфікаційна робота бакалавра присвячена розробці веб-додатку PlantShop для продажу рослин та представлення послуг ландшафтного дизайну з використанням сучасних технологій Next.js та NestJS. У роботі розглянуто предметну область онлайн-продажу рослин, визначено вимоги до системи, описано архітектуру full-stack застосунку та реалізацію його основних модулів.",
        "Робота складається зі вступу, трьох розділів, висновків, списку використаних джерел і додатків. У першому розділі досліджено особливості електронної комерції у сфері продажу рослин, обґрунтовано потребу в цифровому каталозі, особистому кабінеті користувача, історії замовлень та адміністративній панелі. У другому розділі сформовано функціональні й нефункціональні вимоги, визначено ролі користувачів, описано маршрути frontend, основні API endpoint-и backend та обґрунтовано вибір технологічного стеку. У третьому розділі наведено практичну реалізацію PlantShop, структуру frontend і backend, базу даних, механізми автентифікації, role-based access, роботу з медіафайлами, site settings і тестування.",
        "У результаті виконання роботи створено веб-додаток, що містить публічну частину з головною сторінкою і каталогом рослин, кошик із локальним збереженням, backend-backed оформлення замовлень для авторизованих користувачів, кабінет користувача, історію замовлень, налаштування профілю та адміністративну панель для керування рослинами, замовленнями, медіафайлами і налаштуваннями сайту. Система не включає оплату онлайн, гостьове збереження замовлень, інтеграцію служб доставки та повноцінну CMS для всіх сторінок, тому ці напрями розглядаються як перспективи розвитку.",
        "Ключові слова: веб-додаток, PlantShop, Next.js, NestJS, React, TypeScript, Prisma, MySQL, автентифікація, JWT, httpOnly cookies, електронна комерція, адмін-панель, замовлення, медіафайли.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)
    doc.add_page_break()


def add_abstract(doc: Document) -> None:
    add_section_heading(doc, "ABSTRACT")
    paragraphs = [
        "The bachelor qualification work is devoted to the development of the PlantShop web application for selling plants and presenting landscape design services using modern Next.js and NestJS technologies. The work analyzes the application domain, defines system requirements, describes the full-stack architecture, and explains the implementation of the main modules.",
        "The work consists of an introduction, three chapters, conclusions, references, and appendices. The first chapter analyzes the specifics of online plant sales, the need for a digital catalog, a user account area, order history, and an administration panel. The second chapter defines functional and non-functional requirements, user roles, frontend routes, backend API endpoints, and the selected technology stack. The third chapter describes the practical implementation of PlantShop, including frontend and backend structure, database design, authentication and session flow, role-based access control, media management, site settings, and testing.",
        "As a result, a web application was developed that includes a public storefront, a plant catalog, a localStorage-based cart, authenticated order creation, a user dashboard, order history, profile settings, and an administration panel for managing plants, orders, media files, and site settings. The implemented system does not include online payment, persisted guest checkout, delivery provider integration, invoices, or a full CMS for all pages; these features are identified as possible future improvements.",
        "Keywords: web application, PlantShop, Next.js, NestJS, React, TypeScript, Prisma, MySQL, authentication, JWT, httpOnly cookies, e-commerce, admin panel, orders, media files.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    add_section_heading(doc, "ЗМІСТ")
    add_toc_field(doc)
    doc.add_page_break()


def add_intro(doc: Document) -> None:
    add_section_heading(doc, "ВСТУП")
    paragraphs = [
        "Актуальність теми. Продаж рослин через інтернет має низку особливостей, які відрізняють його від типового продажу стандартизованих товарів. Покупцю потрібно не лише побачити назву й ціну, а й оцінити зовнішній вигляд рослини, наявність фотографій, короткий опис, можливість додати товар у кошик і зрозуміти порядок оформлення замовлення. Для бізнесу, який поєднує продаж рослин із послугами ландшафтного дизайну, веб-додаток також виконує роль цифрової вітрини, каналу комунікації та інструмента адміністрування контенту.",
        "Традиційний підхід, за якого каталог ведеться вручну в соціальних мережах або таблицях, не забезпечує достатньої структурованості даних. Адміністратору складно підтримувати актуальні фото, ціни, описи та контактну інформацію, а користувач не отримує стабільного сценарію перегляду товарів і повторного доступу до історії своїх покупок. Тому розробка окремого веб-додатку для такого бізнесу є доцільною, якщо він реалізує не декоративний сайт-візитку, а пов’язану систему з frontend, backend, базою даних, автентифікацією та адміністративними інструментами.",
        "У межах цієї роботи розглядається веб-додаток PlantShop. Його реалізація побудована як full-stack система: frontend-частина створена на Next.js з використанням React і TypeScript, backend-частина реалізована на NestJS, а збереження даних виконується у MySQL через Prisma ORM. Такий підхід дозволяє розділити презентаційний рівень, бізнес-логіку та рівень доступу до даних, не змішуючи їх в одному файлі або одному клієнтському стані.",
        "Метою роботи є розробка та технічний опис веб-додатку PlantShop для продажу рослин і представлення послуг ландшафтного дизайну, який забезпечує публічний каталог, автентифікацію, оформлення замовлень авторизованими користувачами, особистий кабінет і адміністративну панель для керування основними даними системи.",
        "Об’єктом дослідження є процес розробки full-stack веб-додатку для електронної комерції у сфері продажу рослин. Предметом дослідження є архітектурні, програмні та організаційні рішення, використані при створенні PlantShop: маршрути frontend, API endpoint-и backend, Prisma-моделі, механізми автентифікації, керування ролями, оформлення замовлень, робота з медіафайлами та тестування.",
        "Для досягнення мети було поставлено такі завдання: проаналізувати предметну область онлайн-продажу рослин; визначити ролі користувачів і вимоги до системи; спроєктувати архітектуру frontend/backend; реалізувати каталог рослин, кошик і сценарій створення замовлення; реалізувати автентифікацію з email verification, password reset і login challenge; створити кабінет користувача з історією замовлень і налаштуваннями; реалізувати адміністративну панель для рослин, замовлень, медіафайлів і налаштувань сайту; описати базу даних; підготувати тестування основних сценаріїв.",
        "Методами розробки є компонентний підхід у frontend-частині, модульна архітектура NestJS у backend-частині, об’єктно-реляційне відображення даних через Prisma, валідація вхідних DTO, розмежування доступу за ролями, cookie-based session flow з JWT і refresh tokens, а також автоматизоване тестування окремих сервісів, контролерів і клієнтських компонентів.",
        "Практичне значення отриманих результатів полягає в тому, що PlantShop може використовуватися як основа для невеликого онлайн-магазину рослин із базовою адміністративною підсистемою. Реалізована система вже підтримує ключові сценарії: перегляд публічного каталогу, роботу з кошиком, створення реального замовлення авторизованим користувачем, перегляд історії замовлень, редагування профілю та адміністрування каталогу, замовлень, медіафайлів і контентних налаштувань.",
        "Робота не описує як реалізовані ті можливості, яких немає в репозиторії: онлайн-оплату, інтеграцію служб доставки, інвойси, гостьове persisted checkout, повноцінну CMS для всіх сторінок, зміну email, видалення акаунта або production cloud storage. Такі можливості можуть розглядатися лише як подальший розвиток системи після стабілізації базових commerce-сценаріїв.",
        "Кваліфікаційна робота складається зі вступу, трьох розділів, висновків, списку використаних джерел і додатків. У першому розділі розглянуто предметну область і постановку завдання. У другому розділі наведено вимоги, ролі, архітектуру та вибір технологій. У третьому розділі описано практичну реалізацію PlantShop, базу даних, захист інформаційно-комп’ютерної системи та тестування.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)


def add_chapter_1(doc: Document) -> None:
    add_chapter_heading(
        doc,
        "РОЗДІЛ 1",
        "АНАЛІЗ ПРЕДМЕТНОЇ ОБЛАСТІ ТА ПОСТАНОВКА ЗАВДАННЯ НА КВАЛІФІКАЦІЙНУ РОБОТУ",
    )
    add_section_heading(
        doc,
        "1.1 Аналіз сучасного стану проблеми онлайн-продажу рослин та послуг ландшафтного дизайну",
    )
    paragraphs = [
        "Онлайн-продаж рослин поєднує характеристики електронної комерції, візуального каталогу та консультаційного сервісу. Покупець обирає не лише товар із фіксованими параметрами, а живий об’єкт, для якого важливі зовнішній стан, розмір, вигляд листя, умови догляду й доречність у конкретному просторі. Тому веб-додаток для такої предметної області має бути більше ніж списком товарів: він повинен допомагати користувачу швидко переглянути асортимент, порівняти варіанти, зберегти вибір у кошику й оформити замовлення з мінімальною кількістю зайвих дій.",
        "Окрема складність полягає в тому, що бізнес із продажу рослин часто паралельно надає послуги озеленення та ландшафтного дизайну. Для таких послуг сайт виконує презентаційну функцію: пояснює напрям діяльності, демонструє підхід до роботи, дає контакти та формує довіру. У PlantShop ця частина представлена головною сторінкою, контактною сторінкою, footer-контентом і налаштуваннями сайту, які адміністратор може змінювати без прямого редагування коду.",
        "Проблема онлайн-продажу рослин полягає не лише у створенні каталогу. Потрібно забезпечити оновлення цін і фотографій, контроль прав доступу, створення замовлень, збереження історії покупок і можливість керування замовленнями з боку адміністратора. Якщо ці процеси залишаються розділеними між месенджерами, таблицями та вручну оновлюваними сторінками, збільшується ризик помилок: товар може бути застарілим, фото може не відповідати актуальному запису, а історія звернень і покупок залишається розпорошеною.",
        "Цифровий каталог у такій системі має дві функції. Для користувача він є основним інтерфейсом вибору рослин. Для адміністратора він є керованим набором записів, де кожна рослина має назву, опис, ціну й масив зображень. У PlantShop каталог працює через публічні endpoint-и `GET /plants` і `GET /plants/:id`, а адміністративні операції створення, редагування та видалення рослин захищені роллю `ADMIN`.",
        "Кошик є проміжним станом між переглядом каталогу та створенням замовлення. У PlantShop він реалізований у frontend-частині та зберігається у `localStorage`, що дозволяє не втрачати вибрані товари під час переходів між сторінками. Водночас сам факт наявності товарів у кошику не створює persisted order у базі даних. Реальне замовлення формується тільки після автентифікації користувача і відправлення даних на backend endpoint `POST /orders`.",
        "Особистий кабінет користувача потрібний для того, щоб покупець мав доступ до власних даних і не сприймав замовлення як одноразову форму без подальшого відстеження. У PlantShop кабінет містить оглядову сторінку, сторінку історії замовлень і сторінку налаштувань профілю. Такий підхід відрізняє реалізовану систему від простого статичного сайту: користувацькі дані зберігаються на backend, а історія замовлень завантажується з бази даних для поточного акаунта.",
        "Історія замовлень має значення і для користувача, і для адміністратора. Користувач бачить лише власні замовлення через endpoint-и `GET /orders/me` та `GET /orders/me/:id`. Адміністратор має окремі endpoint-и `GET /orders/admin`, `GET /orders/admin/:id` та `PATCH /orders/admin/:id`, які дозволяють переглядати всі замовлення і змінювати їхній статус. Це забезпечує базовий адміністративний workflow без створення надлишкової аналітичної системи, якої в репозиторії немає.",
        "Адміністративна панель у PlantShop виконує роль робочого інструмента для оператора системи. Вона не є окремим експериментальним розділом, а входить до маршруту `/admin` і має власний layout через `AdminShell`. Адміністратор може працювати з каталогом, замовленнями, медіатекою та налаштуваннями сайту. Права доступу перевіряються не тільки на рівні UI, а й на backend через `JwtAuthGuard` і `RolesGuard`.",
        "Керування медіафайлами є важливим саме для предметної області рослин, оскільки фотографії є ключовим фактором вибору. У PlantShop backend зберігає керовані файли у папках `uploads/plants` і `uploads/site`, а публічні URL формуються через backend public base URL. Система перевіряє тип файлу, розмір і колекцію, а видалення файлів враховує, чи використовується зображення хоча б однією рослиною.",
        "Керування контентом сайту реалізовано не як повноцінна CMS, а як окремий singleton `SiteSettings`. Такий обсяг відповідає фактичним потребам PlantShop: адміністратор може змінювати назву сайту, hero-тексти, CTA, контактну інформацію, footer-опис і посилання на соціальні мережі. Це зменшує залежність від редагування вихідного коду, але не створює універсальної системи сторінок, блоків і шаблонів.",
        "Послуги ландшафтного дизайну в контексті PlantShop подані через інформаційні сторінки і маркетингові блоки. У репозиторії є сторінки `/landscape-design`, `/care`, `/pots`, `/contact`, які підтримують контентну частину бізнесу. Водночас окремого повного workflow замовлення консультації, маршруту `/consultation` або портфоліо `/portfolio` у репозиторії не знайдено, тому ці можливості не можна описувати як реалізовані.",
        "Доцільність створення власного веб-додатку PlantShop пояснюється тим, що типовий конструктор або сторінка в соціальній мережі не дає достатнього контролю над session flow, ролями, власною базою даних, історією замовлень і адміністративними операціями. Реалізована система дозволяє не тільки показати рослини, а й зв’язати каталог, користувача, замовлення, медіафайли та налаштування сайту в одну технічну структуру.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_table(
        doc,
        "Таблиця 1.1 – Порівняння існуючих підходів до організації онлайн-продажу рослин",
        ["Підхід", "Переваги", "Недоліки", "Висновок для PlantShop"],
        [
            [
                "Соціальні мережі та месенджери",
                "Швидкий старт, проста комунікація, не потрібна складна інфраструктура",
                "Немає структурованої бази товарів, складно підтримувати актуальні ціни та історію замовлень",
                "Підходить для старту, але не замінює власний каталог і кабінет користувача",
            ],
            [
                "Статичний сайт-візитка",
                "Можна презентувати бренд, контакти і послуги",
                "Немає реальних замовлень, ролей, адмін-панелі та persisted data",
                "Недостатньо для задачі продажу рослин з backend-backed orders",
            ],
            [
                "Готовий SaaS-магазин",
                "Є базова e-commerce функціональність, шаблони та оплата",
                "Обмежена гнучкість, залежність від платформи, складніше адаптувати auth/admin workflow",
                "Може бути альтернативою, але не демонструє власну full-stack реалізацію",
            ],
            [
                "Власний full-stack веб-додаток",
                "Контроль над frontend, backend, БД, auth, ролями, замовленнями та медіафайлами",
                "Потребує розробки, тестування і подальшої підтримки",
                "Саме цей підхід обрано для PlantShop",
            ],
        ],
        widths=[4.0, 4.2, 4.2, 4.2],
    )

    add_section_heading(doc, "1.2 Постановка завдання на кваліфікаційну роботу бакалавра")
    paragraphs = [
        "Завдання кваліфікаційної роботи полягає у створенні веб-додатку PlantShop, який відповідає реальним потребам невеликого онлайн-магазину рослин і не підміняє відсутні можливості декларативним описом. Система повинна мати публічний storefront, каталог рослин, кошик, автентифікацію, кабінет користувача, історію замовлень і адміністративну панель.",
        "Публічна частина повинна надавати користувачу доступ до головної сторінки, каталогу рослин, інформаційних сторінок, контактів і базового перегляду товарів. Каталог має завантажуватися з backend, а не бути лише локальним масивом у компоненті. Перегляд деталі рослини реалізується через модальне вікно, оскільки окремого маршруту `/plants/[id]` у репозиторії не знайдено.",
        "Користувацька частина повинна підтримувати реєстрацію, підтвердження email, вхід, logout, відновлення пароля та login challenge для невпізнаного контексту. Після автентифікації користувач отримує доступ до dashboard, історії замовлень і налаштувань профілю. Тимчасова недоступність auth backend не повинна трактуватися як підтверджений вихід з акаунта.",
        "Сценарій замовлення повинен бути реалізований без guest checkout. Кошик може існувати локально, але persisted order створюється тільки після підтвердження authenticated user. Під час створення замовлення email покупця береться з акаунта, а не з довільного поля checkout-форми.",
        "Адміністративна частина повинна бути доступна лише користувачу з роллю `ADMIN`. Адміністратор має керувати рослинами, замовленнями, медіафайлами і site settings. Для кожної адміністративної дії backend має перевіряти роль, оскільки frontend-обмеження не є достатнім механізмом безпеки.",
        "Backend має бути реалізований як модульний NestJS API. Основними модулями є `auth`, `plants`, `orders`, `media`, `site-settings` і `prisma`. Дані мають зберігатися у MySQL через Prisma. Схема бази даних повинна відображати користувачів, рослини, замовлення, позиції замовлень, refresh tokens, verification tokens, login challenges, trusted login contexts, pending registrations і site settings.",
        "Окреме завдання стосується захисту системи. Необхідно використати httpOnly cookies, JWT access tokens, refresh tokens, bcrypt для хешування, DTO validation, reCAPTCHA для login/register, email verification, password reset і role-based access. Файли uploads повинні зберігатися в керованих папках і не видалятися, якщо вони ще використовуються рослинами.",
        "Результатом виконання завдання має стати не універсальна e-commerce платформа, а обмежений, але цілісний веб-додаток PlantShop. У перспективи розвитку можуть бути винесені online payment, інтеграція служб доставки, інвойси, гостьове persisted checkout, повноцінна CMS, export data, видалення акаунта, зміна email і cloud storage для медіафайлів.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_section_heading(doc, "Висновки до розділу 1")
    conclusions = [
        "У першому розділі визначено, що предметна область онлайн-продажу рослин потребує не лише публічного каталогу, а й підтримки користувацького кабінету, історії замовлень, адміністративного керування даними, медіафайлами і базовими налаштуваннями сайту.",
        "Встановлено, що власний full-stack веб-додаток є доцільним для PlantShop, оскільки дозволяє контролювати структуру даних, auth/session flow, ролі користувачів, backend-backed orders і адміністративні операції. Також сформульовано межі реалізації, щоб не описувати як готові можливості ті функції, яких немає в репозиторії.",
    ]
    for paragraph in conclusions:
        add_plain_paragraph(doc, paragraph)


def add_chapter_2(doc: Document) -> None:
    add_chapter_heading(
        doc,
        "РОЗДІЛ 2",
        "СПЕЦИФІКАЦІЯ ВИМОГ ДО ВЕБ-ДОДАТКУ PLANTSHOP",
    )
    add_section_heading(doc, "2.1 Аналіз, визначення вимог та проєктування програмного забезпечення")
    paragraphs = [
        "Веб-додаток PlantShop має три основні категорії користувачів: anonymous, USER і ADMIN. Anonymous-користувач може переглядати публічні сторінки, відкривати каталог рослин, додавати товари до локального кошика, реєструватися, входити в акаунт, підтверджувати email і запускати сценарій відновлення пароля. USER отримує доступ до backend-backed сценаріїв: створення замовлення, перегляду кабінету, історії замовлень і налаштувань профілю. ADMIN має всі можливості USER і додатково працює з адмін-панеллю.",
        "Поділ ролей у системі є мінімальним і відповідає фактичній Prisma-схемі: enum `UserRole` містить лише `USER` і `ADMIN`. У роботі не вводиться розширена RBAC-матриця з ролями менеджера, оператора доставки або модератора, оскільки таких ролей у репозиторії немає. Доступ до адміністративних frontend routes перевіряється у `frontend/proxy.ts`, а доступ до backend write endpoint-ів додатково перевіряється через `JwtAuthGuard` і `RolesGuard`.",
        "Функціональні вимоги визначено на основі реалізованих маршрутів frontend, backend controllers і тестів. Вони охоплюють публічний перегляд сайту, каталог, auth flow, кабінет користувача, замовлення і адмін-підсистему. Нефункціональні вимоги стосуються безпеки, надійності, підтримуваності, конфігурованості й тестованості.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_table(
        doc,
        "Таблиця 2.1 – Функціональні вимоги до веб-додатку PlantShop",
        ["ID", "Роль", "Вимога", "Реалізація"],
        [
            ["F01", "anonymous", "Перегляд головної сторінки", "`frontend/app/page.tsx`, `SiteSettingsContext`"],
            ["F02", "anonymous", "Перегляд каталогу рослин", "`GET /plants`, `frontend/app/plants/page.tsx`"],
            ["F03", "anonymous", "Робота з кошиком", "`CartContext`, `CartModal`, `localStorage`"],
            ["F04", "anonymous", "Реєстрація та підтвердження email", "`POST /auth/register`, `POST /auth/verify-email`"],
            ["F05", "USER", "Створення замовлення", "`POST /orders`, `OrdersService.createForUser`"],
            ["F06", "USER", "Перегляд історії замовлень", "`GET /orders/me`, `/dashboard/orders`"],
            ["F07", "USER", "Редагування профілю", "`PATCH /auth/me`, `/dashboard/settings`"],
            ["F08", "ADMIN", "Керування рослинами", "`POST/PATCH/DELETE /plants`, `/admin/plants`"],
            ["F09", "ADMIN", "Керування замовленнями", "`GET/PATCH /orders/admin`, `/admin/orders`"],
            ["F10", "ADMIN", "Керування медіа та site settings", "`/admin/media`, `/site-settings`"],
        ],
        widths=[1.2, 2.2, 5.0, 7.6],
    )

    add_table(
        doc,
        "Таблиця 2.2 – Нефункціональні вимоги до веб-додатку PlantShop",
        ["Категорія", "Вимога", "Реалізація", "Перевірка"],
        [
            ["Безпека", "Захищені сесії", "httpOnly cookies, JWT, refresh tokens", "Login/logout, headers"],
            ["Безпека", "Розмежування ролей", "`USER`, `ADMIN`, `RolesGuard`", "Backend tests, manual `/admin`"],
            ["Надійність", "Не фейкувати logout", "state очищається після backend confirmation", "`AuthContext.test.tsx`"],
            ["Надійність", "Не фейкувати guest checkout", "anonymous спрямовується в auth flow", "`CartModal.test.tsx`"],
            ["Підтримуваність", "Модульна структура", "NestJS modules", "Code review"],
            ["Конфігурація", "Env-driven URLs", "`.env.example`, `env.ts`", "Manual env check"],
            ["Тестованість", "Unit та integration tests", "Jest, Vitest", "`npm test`"],
            ["E2E", "Browser smoke scenarios", "Playwright config", "Потребує ручної перевірки"],
        ],
        widths=[2.6, 4.0, 5.2, 4.2],
    )

    add_plain_paragraph(
        doc,
        "Ролі користувачів та їхні права доступу наведено в таблиці 2.3. Таблиця відображає фактичний стан репозиторію і не містить ролей, яких немає у Prisma enum `UserRole`.",
    )
    add_table(
        doc,
        "Таблиця 2.3 – Ролі користувачів та права доступу",
        ["Роль", "Frontend", "Backend API", "Обмеження"],
        [
            [
                "anonymous",
                "`/`, `/plants`, `/about`, `/contact`, auth/security routes",
                "`GET /plants`, `GET /site-settings/public`, public `/auth/*`",
                "Не має persisted checkout, dashboard і admin",
            ],
            [
                "USER",
                "`/dashboard`, `/dashboard/orders`, `/dashboard/settings`, `/logout`",
                "`/auth/me`, `/orders`, `/orders/me`",
                "Не має admin endpoint-ів",
            ],
            [
                "ADMIN",
                "`/admin`, `/admin/plants`, `/admin/orders`, `/admin/media`, `/admin/settings`",
                "Admin writes for plants, orders, media, site settings",
                "Роль призначається окремим operator workflow",
            ],
        ],
        widths=[2.0, 5.0, 5.4, 3.6],
    )

    add_plain_paragraph(
        doc,
        "Діаграму варіантів використання веб-додатку PlantShop наведено на рисунку 2.1. Вона показує, що ADMIN розширює можливості USER, а USER розширює можливості anonymous-користувача.",
    )
    add_figure_placeholder(doc, "2.1", "Діаграма варіантів використання веб-додатку PlantShop")

    paragraphs = [
        "Frontend маршрути побудовані за правилами Next.js App Router. Публічні сторінки знаходяться у `frontend/app`, зокрема `/`, `/plants`, `/about`, `/contact`, `/care`, `/pots`, `/landscape-design`, `/convert`. Захищені користувацькі сторінки розміщені під `/dashboard`, а адміністративні сторінки - під `/admin`. Legacy entrypoints `/signin`, `/register`, `/forgot-password` і `/verify-login-code` редіректять у native auth modal flow.",
        "Backend API складається з контролерів `AuthController`, `PlantController`, `OrdersController`, `AdminOrdersController`, `MediaController` і `SiteSettingsController`. Частина endpoint-ів є публічною, частина потребує authenticated session, а адміністративні endpoint-и потребують ролі `ADMIN`. Основні endpoint-и наведено у таблиці 2.5.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_table(
        doc,
        "Таблиця 2.4 – Технологічний стек веб-додатку PlantShop",
        ["Рівень", "Технологія", "Версія", "Призначення"],
        [
            ["Frontend", "Next.js", "^16.2.4", "App Router, SSR/ISR, protected routes"],
            ["Frontend", "React", "^19.2.5", "Компонентний UI"],
            ["Frontend", "TypeScript", "^5", "Типізація frontend-коду"],
            ["Frontend", "Tailwind CSS", "^4.1.16", "Адаптивне оформлення UI"],
            ["Frontend", "Radix UI", "^1-2", "Dialog, tabs, select, switch"],
            ["Backend", "NestJS", "^11.0.1", "Модульний API"],
            ["Backend", "Prisma", "^6.19.2", "ORM і доступ до MySQL"],
            ["Backend", "MySQL", "через Prisma", "Збереження users, plants, orders, settings"],
            ["Security", "JWT, bcrypt, Nodemailer, reCAPTCHA", "package.json", "Auth/session і security emails"],
            ["Testing", "Jest, Vitest, Testing Library", "package.json", "Unit та integration testing"],
        ],
        widths=[2.4, 4.0, 2.7, 6.9],
    )

    add_table(
        doc,
        "Таблиця 2.5 – Основні API endpoint-и системи",
        ["Метод", "Endpoint", "Guard/role", "Опис"],
        [
            ["POST", "/auth/register", "public", "Створення pending registration"],
            ["POST", "/auth/login", "public", "Вхід або login challenge"],
            ["GET", "/auth/me", "cookie session", "Поточний користувач"],
            ["PATCH", "/auth/me", "cookie session", "Оновлення профілю"],
            ["GET", "/plants", "public", "Каталог рослин"],
            ["POST", "/plants", "ADMIN", "Створення рослини"],
            ["POST", "/orders", "USER/ADMIN", "Створення замовлення поточного користувача"],
            ["GET", "/orders/me", "USER/ADMIN", "Власні замовлення"],
            ["GET", "/orders/admin", "ADMIN", "Усі замовлення"],
            ["PATCH", "/orders/admin/:id", "ADMIN", "Оновлення статусу"],
            ["POST", "/admin/media/upload", "ADMIN", "Завантаження медіафайлів"],
            ["PATCH", "/site-settings", "ADMIN", "Оновлення налаштувань сайту"],
        ],
        widths=[1.8, 4.8, 3.0, 6.4],
    )

    add_plain_paragraph(
        doc,
        "Загальну архітектуру веб-додатку PlantShop наведено на рисунку 2.2. Вона відображає взаємодію браузера, Next.js frontend, NestJS API, Prisma, MySQL, SMTP, reCAPTCHA та файлового сховища uploads.",
    )
    add_figure_placeholder(doc, "2.2", "Загальна архітектура веб-додатку PlantShop")
    add_plain_paragraph(
        doc,
        "Схему API-взаємодії frontend і backend наведено на рисунку 2.3. У ній доцільно показати окремо публічні запити каталогу, authenticated запити користувача та admin-only операції.",
    )
    add_figure_placeholder(doc, "2.3", "Схема API-взаємодії frontend і backend")

    add_section_heading(doc, "2.2 Вибір засобів, методів і технологій реалізації веб-додатку")
    paragraphs = [
        "Next.js обрано для frontend-частини, оскільки він дозволяє поєднувати серверне завантаження даних, клієнтську інтерактивність і файлову структуру маршрутів. У PlantShop це використано для головної сторінки, каталогу, dashboard, auth routes і admin routes. Для каталогу застосовано server-side fetch до backend API, а компоненти на кшталт `PlantsList`, `CartModal` і `AuthModal` працюють як client components.",
        "React забезпечує компонентну модель інтерфейсу. У PlantShop окремими компонентами винесено навігацію, footer, картки рослин, модальні вікна, admin shell і managers. TypeScript використано для типізації користувача, рослини, замовлення, site settings і API payloads. Це зменшує ризик помилок під час роботи з даними, які приходять з backend.",
        "Tailwind CSS застосовано для швидкого створення адаптивного UI. У репозиторії також використовуються Radix UI компоненти для dialog, tabs, table, select, switch і споріднених елементів. Framer Motion використовується для анімацій modal і storefront-елементів, але анімації не замінюють основну функціональність системи.",
        "NestJS обрано для backend API через модульну структуру і чіткий поділ контролерів, сервісів, DTO та guard-ів. У PlantShop це видно за окремими runtime modules: `auth`, `plants`, `orders`, `media`, `site-settings`, `prisma`. Така структура полегшує тестування і локалізує зміни в межах відповідного домену.",
        "Prisma використовується як ORM для MySQL. Схема `backend/prisma/schema.prisma` містить моделі, які відповідають реальним сутностям системи: користувачі, рослини, замовлення, позиції замовлень, токени, login challenges, trusted contexts, pending registrations і site settings. Це дозволяє описувати базу даних декларативно і підтримувати типізований доступ у backend-сервісах.",
        "Для автентифікації використано JWT у поєднанні з httpOnly cookies. Access token має короткий строк дії, refresh token зберігається у БД у вигляді hash, а trusted login context дозволяє розпізнавати браузерний контекст. bcrypt використано для хешування паролів і порівняння секретних значень. Nodemailer відповідає за email verification, password reset і login challenge, а reCAPTCHA використовується для login/register flow.",
        "Для тестування backend використовується Jest. Тести покривають DTO validation, auth service/controller, roles guard, plants, orders, media і site settings. Для frontend використовується Vitest, Testing Library і happy-dom. Playwright config наявний для browser smoke scenarios з mock backend, однак фактична доступність Playwright tooling потребує ручної перевірки, оскільки у `frontend/package.json` Playwright не зафіксовано як локальну devDependency.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_section_heading(doc, "Висновки до розділу 2")
    conclusions = [
        "У другому розділі визначено ролі користувачів, функціональні й нефункціональні вимоги, frontend routes, backend API endpoint-и та загальну архітектуру PlantShop. Система має чіткий поділ між публічними, authenticated і admin-only сценаріями.",
        "Обґрунтовано використання Next.js, React, TypeScript і Tailwind CSS для frontend, NestJS і Prisma для backend, MySQL для збереження даних, а також JWT, httpOnly cookies, bcrypt, Nodemailer і reCAPTCHA для auth/security flow. Вибір технологій відповідає фактичному коду репозиторію.",
    ]
    for paragraph in conclusions:
        add_plain_paragraph(doc, paragraph)


def add_chapter_3(doc: Document) -> None:
    add_chapter_heading(
        doc,
        "РОЗДІЛ 3",
        "РОЗРОБКА ВЕБ-ДОДАТКУ PLANTSHOP ДЛЯ ПРОДАЖУ РОСЛИН ТА НАДАННЯ ПОСЛУГ ЛАНДШАФТНОГО ДИЗАЙНУ",
    )
    add_section_heading(doc, "3.1 Практична реалізація об’єкта проєктування")
    paragraphs = [
        "Frontend-частина PlantShop розміщена у папці `frontend`. Основний застосунок знаходиться в `frontend/app`, що відповідає структурі Next.js App Router. Публічні сторінки включають головну сторінку, каталог рослин, сторінки про проєкт, контакти, догляд, горщики, ландшафтний дизайн і конвертер. Захищені користувацькі сторінки розміщені в `frontend/app/dashboard`, а адміністративні сторінки - в `frontend/app/admin`.",
        "Глобальна структура frontend формується через `frontend/app/layout.tsx` і `frontend/app/ClientWrapper.tsx`. Root layout підключає глобальні стилі, metadata, viewport settings і передає children у client wrapper. `ClientWrapper` підключає провайдери теми, notification system, site settings, auth, cart і wishlist. Для admin routes звичайний navbar і footer не рендеряться, оскільки адміністративна частина має власний shell.",
        "Контексти frontend виконують роль локального стану й інтеграційного шару з backend. `AuthContext` відповідає за bootstrap поточного користувача через `/auth/me`, login, register, login challenge, logout і redirect у login flow. `CartContext` зберігає кошик у `localStorage`, `WishlistContext` відповідає за обрані товари, а `SiteSettingsContext` завантажує публічні налаштування сайту через `/site-settings/public`.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_plain_paragraph(
        doc,
        "Фрагмент підключення frontend-провайдерів наведено у лістингу 3.1. Він показує, що auth, cart, wishlist і site settings є глобальними частинами клієнтського застосунку.",
    )
    add_listing(
        doc,
        "3.1",
        "Підключення основних frontend context providers",
        """
<SiteSettingsProvider>
  <AuthProvider>
    <CartProvider>
      <WishlistProvider>
        <EnhancedThemeRipple />
        <PWAInstaller />
        {children}
      </WishlistProvider>
    </CartProvider>
  </AuthProvider>
</SiteSettingsProvider>
""",
    )

    paragraphs = [
        "Каталог рослин реалізовано сторінкою `frontend/app/plants/page.tsx` і компонентом `PlantsList`. Сторінка виконує server-side fetch до backend `GET /plants`, а компонент відповідає за відображення карток, модальний перегляд, додавання до кошика і wishlist. Для адміністратора в admin context також доступні дії створення, редагування і видалення рослин, однак основний admin workflow винесено в `PlantsAdminManager`.",
        "Окремої сторінки `/plants/[id]` у репозиторії не знайдено. Тому в роботі деталізацію рослини потрібно описувати саме як модальне вікно у каталозі. Модальне вікно показує зображення, назву, опис, ціну, додаткові характеристики UI та кнопки додавання до кошика. Частина текстів у UI має маркетинговий характер, але persisted product data надходить з backend моделі `Plant`.",
        "Кошик реалізовано у `CartContext` і `CartModal`. Кошик може існувати для anonymous-користувача, але persisted order створюється тільки після входу. Якщо користувач не авторизований, `CartModal` не симулює покупку, а закриває модальне вікно і спрямовує користувача в auth flow. Це відповідає фактичному продуктовому рішенню: guest checkout не реалізовано.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_plain_paragraph(
        doc,
        "Фрагмент логіки checkout наведено у лістингу 3.2. Він підтверджує, що створення замовлення відбувається тільки для користувача з активною сесією.",
    )
    add_listing(
        doc,
        "3.2",
        "Перевірка авторизації перед створенням замовлення",
        """
if (!user) {
  addNotification({
    type: 'info',
    title: 'Потрібен акаунт',
    message: 'Щоб оформити реальне замовлення, спершу увійдіть.',
  });
  onClose();
  await ensureAuthAndGo('/dashboard/orders');
  return;
}

const order = await createCurrentUserOrder({
  customerName: checkout.customerName.trim(),
  customerPhone: checkout.customerPhone.trim(),
  deliveryAddress: checkout.deliveryAddress.trim(),
  note: checkout.note.trim() || undefined,
  items: orderItems,
});
""",
    )

    paragraphs = [
        "Dashboard користувача реалізовано у `frontend/app/dashboard/page.tsx`. Він завантажує реальну історію замовлень через `listCurrentUserOrders` і показує останнє замовлення, якщо воно існує. У dashboard спеціально не використовуються вигадані метрики продажів або фейкова історія покупок.",
        "Сторінка історії замовлень `frontend/app/dashboard/orders/page.tsx` працює зі списком замовлень і деталями вибраного запису. Deep-link route `frontend/app/dashboard/orders/[id]/page.tsx` не рендерить окрему сторінку деталей, а перенаправляє користувача на канонічний стан `/dashboard/orders?orderId=...`. Це важливо коректно описати в роботі.",
        "Сторінка налаштувань `frontend/app/dashboard/settings/page.tsx` дозволяє змінювати профіль, notification preferences, privacy settings і пароль. Оновлення профілю виконується через `PATCH /auth/me`, а зміна пароля - через backend-backed security flow. На цій сторінці також є disabled дії export data і delete account, але вони не підключені до backend, тому їх потрібно описувати як не реалізовані.",
        "Адмін-панель розміщена під `/admin`. Layout `frontend/app/admin/layout.tsx` підключає `AdminShell`, який містить sidebar navigation до огляду, рослин, замовлень, медіа і налаштувань. Сторінка `/admin` показує короткі дані про каталог і site settings, а робочі сторінки делегують функціональність менеджерам у `frontend/components/admin`.",
        "`PlantsAdminManager` дозволяє адміністратору створювати, редагувати і видаляти рослини. `OrdersAdminManager` завантажує список замовлень і дозволяє змінювати статус. `MediaManager` і `MediaUploader` працюють із upload endpoint-ами backend. `SiteSettingsManager` редагує singleton налаштувань сайту і після збереження оновлює frontend context.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_figure_placeholder(doc, "3.4", "Головна сторінка веб-додатку PlantShop")
    add_figure_placeholder(doc, "3.5", "Каталог рослин")
    add_figure_placeholder(doc, "3.6", "Модальне вікно перегляду рослини")
    add_figure_placeholder(doc, "3.7", "Кошик користувача")
    add_figure_placeholder(doc, "3.8", "Особистий кабінет користувача")
    add_figure_placeholder(doc, "3.9", "Історія замовлень користувача")
    add_figure_placeholder(doc, "3.10", "Панель адміністратора")
    add_figure_placeholder(doc, "3.11", "Керування рослинами в адмін-панелі")
    add_figure_placeholder(doc, "3.12", "Керування замовленнями в адмін-панелі")
    add_figure_placeholder(doc, "3.13", "Керування медіафайлами")
    add_figure_placeholder(doc, "3.14", "Налаштування сайту")

    add_plain_paragraph(
        doc,
        "Backend-частина розміщена у папці `backend`. Точкою входу є `backend/src/main.ts`, а кореневим модулем - `backend/src/app.module.ts`. У процесі запуску backend перевіряє обов’язкові env-змінні, створює upload directories, підключає cookie parser, статичну роздачу `/uploads`, глобальний `ValidationPipe` і CORS з credentials.",
    )
    add_plain_paragraph(
        doc,
        "Фрагмент bootstrap backend API наведено у лістингу 3.3. Він демонструє конфігурацію cookies, uploads, DTO validation і CORS.",
    )
    add_listing(
        doc,
        "3.3",
        "Bootstrap NestJS API",
        """
const app = await NestFactory.create(AppModule);

app.use(cookieParser());
app.use('/uploads', expressStatic(getUploadsRootDir()));

app.useGlobalPipes(
  new ValidationPipe({
    whitelist: true,
    transform: true,
    transformOptions: { enableImplicitConversion: true },
  }),
);

app.enableCors({
  origin: corsOrigin,
  credentials: true,
});
""",
    )

    add_plain_paragraph(
        doc,
        "Склад основних backend-модулів наведено у лістингу 3.4. Такий поділ дозволяє відокремити auth, каталог, замовлення, медіафайли, налаштування сайту і доступ до бази даних.",
    )
    add_listing(
        doc,
        "3.4",
        "Підключення модулів NestJS",
        """
@Module({
  imports: [
    ThrottlerModule.forRoot({ throttlers: [{ ttl: 60_000, limit: 1000 }] }),
    AuthModule,
    PrismaModule,
    PlantsModule,
    MediaModule,
    SiteSettingsModule,
    OrdersModule,
  ],
})
export class AppModule {}
""",
    )

    add_section_heading(doc, "3.2 Розробка бази даних")
    paragraphs = [
        "База даних PlantShop описана у `backend/prisma/schema.prisma`. Як datasource використовується MySQL, а Prisma Client генерується на основі декларативної схеми. Основні домени бази даних: каталог рослин, користувачі, security tokens, site settings і замовлення.",
        "Модель `Plant` містить назву, опис, ціну, часові мітки та поле `images`, яке зберігається як JSON-serialized string у текстовому полі. Такий підхід уже використовується сервісом `PlantsService`: під час запису images серіалізуються, а під час відповіді перетворюються назад у масив рядків.",
        "Модель `User` зберігає email, password hash, дату підтвердження email, профільні поля, notification preferences, privacy settings і роль. З користувачем пов’язані refresh tokens, trusted login contexts, verification tokens, login challenges і orders. Нові користувачі мають роль `USER` за замовчуванням.",
        "Orders domain складається з моделей `Order` і `OrderItem`. `Order` належить користувачу і містить статус, контактні поля, адресу доставки, коментар і загальну суму. `OrderItem` зберігає snapshot позиції: назву рослини, зображення, ціну за одиницю, кількість і line total. Це важливо, тому що після зміни рослини в каталозі історичне замовлення має залишатися зрозумілим.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_table(
        doc,
        "Таблиця 3.1 – Основні моделі бази даних PlantShop",
        ["Модель", "Ключові поля", "Зв’язки", "Призначення"],
        [
            ["Plant", "`id`, `name`, `description`, `price`, `images`", "`OrderItem[]`", "Каталог рослин"],
            ["User", "`email`, `password`, `emailVerifiedAt`, `role`", "tokens, challenges, orders", "Акаунти користувачів"],
            ["PendingRegistration", "`email`, `passwordHash`, `tokenHash`", "standalone", "Реєстрація до verify email"],
            ["RefreshToken", "`tokenHash`, `expiresAt`, `userId`", "`User`", "Refresh sessions"],
            ["TrustedLoginContext", "`tokenHash`, `userAgentSummary`", "`User`", "Розпізнаний браузерний контекст"],
            ["VerificationToken", "`type`, `tokenHash`, `expiresAt`", "`User`", "Email verify і password reset"],
            ["LoginChallenge", "`codeHash`, `contextHash`, `attemptCount`", "`User`", "Код підтвердження входу"],
            ["SiteSettings", "`siteName`, `hero*`, `contact*`, `footer*`", "singleton", "Контентні налаштування"],
            ["Order", "`userId`, `status`, `totalAmount`", "`User`, `OrderItem[]`", "Замовлення"],
            ["OrderItem", "`plantName`, `unitPrice`, `quantity`", "`Order`, optional `Plant`", "Позиція замовлення"],
        ],
        widths=[3.0, 4.8, 3.6, 4.6],
    )

    add_table(
        doc,
        "Таблиця 3.2 – Статуси замовлення",
        ["Статус", "Значення в enum", "Призначення"],
        [
            ["Нове", "PENDING", "Замовлення створено, але ще не обробляється"],
            ["В обробці", "PROCESSING", "Адміністратор почав роботу із замовленням"],
            ["Відправлено", "SHIPPED", "Замовлення передано на етап доставки"],
            ["Доставлено", "DELIVERED", "Замовлення завершено"],
            ["Скасовано", "CANCELLED", "Замовлення скасовано"],
        ],
        widths=[4.0, 4.0, 8.0],
    )

    add_plain_paragraph(
        doc,
        "Логічну структуру бази даних PlantShop наведено на рисунку 3.1. На діаграмі доцільно показати зв’язки `User` з `Order`, `Order` з `OrderItem`, `Plant` з `OrderItem`, а також security-моделі користувача.",
    )
    add_figure_placeholder(doc, "3.1", "Логічна структура бази даних PlantShop")

    add_plain_paragraph(
        doc,
        "Фрагмент Prisma-схеми для домену замовлень наведено у лістингу 3.5. Він показує статуси замовлень і зв’язок між замовленням та його позиціями.",
    )
    add_listing(
        doc,
        "3.5",
        "Моделі замовлення у Prisma schema",
        """
enum OrderStatus {
  PENDING
  PROCESSING
  SHIPPED
  DELIVERED
  CANCELLED
}

model Order {
  id          Int         @id @default(autoincrement())
  userId      Int
  user        User        @relation(fields: [userId], references: [id])
  status      OrderStatus @default(PENDING)
  totalAmount Float
  items       OrderItem[]
}
""",
    )

    add_plain_paragraph(
        doc,
        "Сервіс створення замовлення отримує authenticated user і DTO з checkout-форми. Він читає рослини з бази даних, формує snapshot позицій, обчислює загальну суму і створює `Order` разом з `OrderItem`. Ключовий фрагмент наведено у лістингу 3.6.",
    )
    add_listing(
        doc,
        "3.6",
        "Створення persisted order",
        """
return tx.order.create({
  data: {
    userId: user.userId,
    status: OrderStatus.PENDING,
    customerName: normalizeRequiredText(dto.customerName),
    customerEmail: user.email,
    customerPhone: normalizeRequiredText(dto.customerPhone),
    deliveryAddress: normalizeRequiredText(dto.deliveryAddress),
    totalAmount,
    items: { create: orderItems },
  },
  select: orderDetailSelect,
});
""",
    )

    add_section_heading(doc, "3.3 Захист інформаційно-комп’ютерної системи")
    paragraphs = [
        "Захист PlantShop побудовано на поєднанні backend-driven session flow, DTO validation, role-based access і обмеження небезпечних операцій. Система не використовує frontend-only auth marker як джерело істини. Поточний стан користувача визначається через `/auth/me`, а cookies мають httpOnly атрибут.",
        "Реєстрація створює `PendingRegistration`, а не одразу `User`. Реальний користувач створюється після підтвердження email. Login для непідтвердженого акаунта не видає cookies, а повертає `verification_required`. Якщо контекст входу не розпізнано, backend створює login challenge і завершує session creation тільки після перевірки коду.",
        "Password reset і dashboard change-password відкликають refresh sessions і trusted login contexts. Звичайний logout відкликає refresh token і очищає session cookies, але не обов’язково видаляє trusted login recognition. Такий підхід відповідає реалізованій логіці backend і frontend.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_table(
        doc,
        "Таблиця 3.3 – Основні механізми захисту системи",
        ["Механізм", "Реалізація", "Файл"],
        [
            ["httpOnly cookies", "`access_token`, `refresh_token`, `trusted_login`", "`auth.controller.ts`"],
            ["JWT strategy", "Passport JWT guard", "`jwt.strategy.ts`, `jwt-auth.guard.ts`"],
            ["Role-based access", "`RolesGuard`, `@Roles('ADMIN')`", "`roles.guard.ts`"],
            ["DTO validation", "`ValidationPipe`, class-validator DTO", "`main.ts`, `dto/*.ts`"],
            ["reCAPTCHA", "login/register verification", "`auth.controller.ts`"],
            ["Email security", "verify email, reset password, challenge", "`email-security.service.ts`"],
            ["Uploads safety", "collection, extension, size, safe path", "`media-storage.ts`, `media.controller.ts`"],
            ["Frontend protected routes", "authenticated / unauthorized / unavailable", "`frontend/proxy.ts`"],
        ],
        widths=[4.0, 6.0, 6.0],
    )

    add_plain_paragraph(
        doc,
        "Схему процесу автентифікації та підтримки сесії наведено на рисунку 3.2. На ній потрібно показати register, email verification, login, login challenge, встановлення cookies і перевірку `/auth/me`.",
    )
    add_figure_placeholder(doc, "3.2", "Схема процесу автентифікації та підтримки сесії")

    add_plain_paragraph(
        doc,
        "Фрагмент role-based access наведено у лістингу 3.7. Він демонструє, що backend перевіряє роль користувача перед виконанням admin-only операцій.",
    )
    add_listing(
        doc,
        "3.7",
        "Перевірка ролей у RolesGuard",
        """
if (!requiredRoles?.length) {
  return true;
}

const user = request.user;
if (!user?.role) {
  throw new ForbiddenException('User role is missing');
}

if (!requiredRoles.includes(user.role)) {
  throw new ForbiddenException('Insufficient permissions');
}
""",
    )

    add_plain_paragraph(
        doc,
        "Frontend proxy також бере участь у захисті route access. Його логіка не зводить тимчасову недоступність auth backend до confirmed logout. Відповідний фрагмент наведено у лістингу 3.8.",
    )
    add_listing(
        doc,
        "3.8",
        "Tri-state policy у frontend proxy",
        """
const PROTECTED_PATHS = {
  "/dashboard": { requireAuth: true },
  "/admin": { requireAuth: true, requireRole: "ADMIN" },
} as const;

if (authResult.status === "unavailable") {
  const unavailableUrl = new URL("/auth-unavailable", request.url);
  unavailableUrl.searchParams.set("next", request.nextUrl.pathname);
  return NextResponse.redirect(unavailableUrl);
}
""",
    )

    add_plain_paragraph(
        doc,
        "Клієнтський auth state також розрізняє authenticated, unauthorized і unavailable. Фрагмент перевірки `/auth/me` наведено у лістингу 3.9.",
    )
    add_listing(
        doc,
        "3.9",
        "Перевірка поточної сесії у AuthContext",
        """
const { res, unauthorized, temporaryFailure } =
  await fetchWithRefresh('/auth/me');

if (unauthorized) {
  setUser(null);
  return 'unauthorized';
}

if (temporaryFailure || !res) {
  return 'unavailable';
}
""",
    )

    add_plain_paragraph(
        doc,
        "Для завантаження медіафайлів backend перевіряє тип, розширення, кількість і розмір файлів. Скорочений фрагмент наведено у лістингу 3.10.",
    )
    add_listing(
        doc,
        "3.10",
        "Обмеження завантаження медіафайлів",
        """
limits: {
  fileSize: 5 * 1024 * 1024,
  files: 10,
},
fileFilter: (_req, file, callback) => {
  const isAllowedMime = file.mimetype.startsWith('image/');
  const allowedExtensions = ['.jpg', '.jpeg', '.png', '.webp', '.gif'];
  callback(null, isAllowedMime && allowedExtensions.includes(extension));
}
""",
    )

    add_plain_paragraph(
        doc,
        "Безпечне перетворення public upload URL у локальний шлях наведено у лістингу 3.11. Цей фрагмент потрібний для пояснення, чому видалення файлів не виконується для довільних шляхів файлової системи.",
    )
    add_listing(
        doc,
        "3.11",
        "Безпечне визначення шляху керованого upload-файлу",
        """
if (!pathname.startsWith('/uploads/')) {
  return null;
}

const relativePath = normalize(pathname.replace(/^\\/uploads\\//, ''));
const filePath = join(UPLOADS_ROOT, relativePath);

if (!filePath.startsWith(UPLOADS_ROOT)) {
  return null;
}
""",
    )

    add_plain_paragraph(
        doc,
        "Сценарій оформлення замовлення наведено на рисунку 3.3. У ньому треба показати каталог, кошик, перевірку authenticated user, `POST /orders`, створення записів `Order` і `OrderItem`, а також перехід до історії замовлень.",
    )
    add_figure_placeholder(doc, "3.3", "Схема процесу оформлення замовлення")

    add_plain_paragraph(
        doc,
        "Site settings реалізовано як singleton-запис у базі даних. Скорочений фрагмент сервісу наведено у лістингу 3.12.",
    )
    add_listing(
        doc,
        "3.12",
        "Отримання singleton налаштувань сайту",
        """
const SETTINGS_ID = 1;

async getSettings() {
  return this.prisma.siteSettings.upsert({
    where: { id: SETTINGS_ID },
    update: {},
    create: { id: SETTINGS_ID, ...defaultSettings },
  });
}
""",
    )

    add_section_heading(doc, "3.4 Тестування та налагодження веб-додатку")
    paragraphs = [
        "Тестування PlantShop організовано на кількох рівнях. Backend має Jest-тести для сервісів, контролерів, DTO validation, role guard, admin contracts, media storage, orders і site settings. Frontend має Vitest і Testing Library тести для контекстів, модальних вікон, dashboard, admin components, auth-security helper-ів і fetchWithRefresh.",
        "Backend-тести підтверджують, що реєстрація створює pending registration, login не видає cookies для непідтвердженого акаунта, login challenge ставить cookies після успішної перевірки, `/auth/me` працює з refresh token, current-user orders не повертають чужі замовлення, а admin-only endpoints вимагають роль `ADMIN`.",
        "Frontend-тести перевіряють, що logout не очищає користувача без підтвердження backend, anonymous checkout не створює фейкове замовлення, dashboard показує реальне останнє замовлення, `/dashboard/orders` працює зі списком і деталями, а admin managers можуть завантажувати й оновлювати дані.",
        "У frontend також наявний Playwright config для browser E2E smoke flows з mock backend. Він перевіряє admin access, admin plants, admin settings і user cabinet smoke flow. Однак фактичний запуск Playwright потребує ручної перевірки середовища, оскільки залежність Playwright не зафіксована як локальна devDependency у `frontend/package.json`.",
        "Backend має opt-in real-stack smoke suite у `backend/test/admin-smoke.e2e-spec.ts`. Він запускається лише за умови `RUN_REAL_STACK_SMOKE=1` і потребує доступної тестової бази даних. Цей тестовий рівень не можна вважати виконаним без ручного запуску з коректним `DATABASE_URL`.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)

    add_table(
        doc,
        "Таблиця 3.4 – План тестування веб-додатку PlantShop",
        ["Рівень", "Що перевіряється", "Файли", "Статус"],
        [
            ["Backend unit/integration", "Auth, orders, plants, media, site settings", "`backend/src/**/*.spec.ts`", "Тести є"],
            ["Backend DTO", "Validation rules для auth/orders/plants", "`auth.dto.spec.ts`, `orders.dto.spec.ts`", "Тести є"],
            ["Backend smoke", "Real cookies, admin settings, orders", "`backend/test/admin-smoke.e2e-spec.ts`", "Opt-in, потребує ручного запуску"],
            ["Frontend component", "Auth modal, cart modal, admin managers", "`frontend/**/*.test.tsx`", "Тести є"],
            ["Frontend helpers", "fetchWithRefresh, auth-security", "`frontend/lib/*.test.ts`", "Тести є"],
            ["Browser E2E", "admin/user smoke flows", "`frontend/e2e`, `playwright.config.ts`", "Потребує ручної перевірки tooling"],
        ],
        widths=[3.0, 5.0, 5.0, 3.0],
    )

    add_section_heading(doc, "Висновки до розділу 3")
    conclusions = [
        "У третьому розділі описано практичну реалізацію PlantShop як full-stack веб-додатку з frontend на Next.js і backend на NestJS. Реалізація охоплює публічний storefront, каталог, кошик, authenticated order creation, dashboard користувача, історію замовлень, профільні налаштування та адмін-панель.",
        "Розроблена база даних відображає основні сутності системи: рослини, користувачі, замовлення, позиції замовлень, security tokens, login challenges, trusted contexts, pending registrations і site settings. Захист системи реалізовано через cookie-based JWT sessions, refresh tokens, bcrypt, email verification, login challenge, DTO validation, CORS і role-based access.",
        "Тестова база репозиторію містить backend Jest-тести, frontend Vitest/Testing Library тести, Playwright config і opt-in real-stack smoke suite. Частина перевірок, пов’язана з реальним середовищем, SMTP, reCAPTCHA, MySQL і Playwright tooling, потребує ручного запуску перед фінальним захистом роботи.",
    ]
    for paragraph in conclusions:
        add_plain_paragraph(doc, paragraph)


def add_conclusions(doc: Document) -> None:
    add_section_heading(doc, "ВИСНОВКИ")
    paragraphs = [
        "У кваліфікаційній роботі бакалавра розроблено та описано веб-додаток PlantShop для продажу рослин і представлення послуг ландшафтного дизайну. Система реалізована як full-stack застосунок із frontend на Next.js і backend на NestJS, а збереження даних виконується у MySQL через Prisma.",
        "У межах роботи проаналізовано предметну область онлайн-продажу рослин. Визначено, що для такого типу бізнесу потрібний не лише публічний сайт, а й керований каталог, кошик, захищена автентифікація, історія замовлень, особистий кабінет і адміністративна панель. Це стало основою для формування функціональних і нефункціональних вимог.",
        "Реалізовано публічну частину PlantShop, яка містить головну сторінку, каталог рослин і модальний перегляд рослини. Кошик зберігається у `localStorage`, але persisted order створюється тільки для авторизованого користувача. Такий підхід не підміняє реальний checkout фейковим сценарієм і відповідає фактичним продуктовим обмеженням системи.",
        "Реалізовано користувацький кабінет, історію замовлень і сторінку налаштувань профілю. Користувач може переглядати власні замовлення через backend-backed endpoint-и, редагувати профільні параметри і змінювати пароль. При цьому видалення акаунта, export data та зміна email не реалізовані й можуть розглядатися тільки як перспективи розвитку.",
        "Реалізовано адміністративну панель для ролі `ADMIN`. Адміністратор може керувати рослинами, переглядати замовлення, змінювати їхній статус, працювати з медіафайлами і редагувати основні налаштування сайту. Доступ до admin routes і backend endpoint-ів обмежено через frontend proxy, JWT guard і role guard.",
        "Розроблена Prisma-схема охоплює моделі `Plant`, `User`, `Order`, `OrderItem`, `SiteSettings`, `RefreshToken`, `VerificationToken`, `LoginChallenge`, `TrustedLoginContext` і `PendingRegistration`. Окрему увагу приділено security flow: email verification, password reset, login challenge, trusted login context, refresh sessions і httpOnly cookies.",
        "Під час роботи підготовлено тестове покриття для backend і frontend. Backend перевіряється Jest-тестами, frontend - Vitest і Testing Library. Playwright config і DB-backed real-stack smoke suite наявні в репозиторії, але їхній фактичний запуск залежить від середовища і потребує ручної перевірки.",
        "Отриманий результат відповідає поставленій меті: створено веб-додаток, який демонструє пов’язану реалізацію storefront, auth, orders, dashboard, admin panel, media management і site settings. Подальший розвиток може включати онлайн-оплату, інтеграцію служб доставки, інвойси, повноцінний workflow консультацій, хмарне зберігання медіафайлів і розширене керування акаунтом.",
    ]
    for paragraph in paragraphs:
        add_plain_paragraph(doc, paragraph)


def add_references(doc: Document) -> None:
    add_section_heading(doc, "СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ")
    references = [
        "Next.js Documentation. URL: https://nextjs.org/docs",
        "React Documentation. URL: https://react.dev",
        "TypeScript Documentation. URL: https://www.typescriptlang.org/docs/",
        "NestJS Documentation. URL: https://docs.nestjs.com",
        "Prisma Documentation. URL: https://www.prisma.io/docs",
        "MySQL Documentation. URL: https://dev.mysql.com/doc/",
        "OWASP Web Security Testing Guide. URL: https://owasp.org/www-project-web-security-testing-guide/",
        "OWASP Cheat Sheet Series. Authentication Cheat Sheet. URL: https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html",
        "JSON Web Tokens. URL: https://jwt.io/introduction",
        "bcrypt package documentation. URL: https://www.npmjs.com/package/bcrypt",
        "Nodemailer Documentation. URL: https://nodemailer.com/about/",
        "Google reCAPTCHA Documentation. URL: https://developers.google.com/recaptcha",
        "Tailwind CSS Documentation. URL: https://tailwindcss.com/docs",
        "Radix UI Documentation. URL: https://www.radix-ui.com/primitives/docs/overview/introduction",
        "Framer Motion Documentation. URL: https://motion.dev/docs",
        "Testing Library Documentation. URL: https://testing-library.com/docs/",
        "Vitest Documentation. URL: https://vitest.dev/guide/",
        "Jest Documentation. URL: https://jestjs.io/docs/getting-started",
        "Playwright Documentation. URL: https://playwright.dev/docs/intro",
        "MDN Web Docs. HTTP Cookies. URL: https://developer.mozilla.org/en-US/docs/Web/HTTP/Cookies",
        "MDN Web Docs. Fetch API. URL: https://developer.mozilla.org/en-US/docs/Web/API/Fetch_API",
        "Practical E-Commerce. Ecommerce UX and product catalog guidance. URL: https://www.practicalecommerce.com/",
        "Nielsen Norman Group. E-Commerce User Experience. URL: https://www.nngroup.com/topic/ecommerce/",
        "Martin Fowler. Patterns of Enterprise Application Architecture. Addison-Wesley, 2002.",
        "Sommerville I. Software Engineering. Pearson, 10th edition, 2015.",
    ]
    for i, ref in enumerate(references, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(f"{i}. {ref}")
        set_run_font(run, 14)


def add_appendices(doc: Document) -> None:
    doc.add_page_break()
    add_section_heading(doc, "ДОДАТКИ")

    add_section_heading(doc, "ДОДАТОК А – Діаграми проєктування веб-додатку PlantShop")
    add_plain_paragraph(
        doc,
        "У додатку А доцільно розмістити повнорозмірні діаграми, які в основному тексті подані у скороченому вигляді: use case diagram, architecture diagram, ERD, auth/session sequence diagram, order flow і media upload flow.",
    )
    for number, title in [
        ("А.1", "Діаграма варіантів використання PlantShop"),
        ("А.2", "Загальна архітектура системи"),
        ("А.3", "ERD бази даних PlantShop"),
        ("А.4", "Auth/session flow"),
        ("А.5", "Order flow"),
        ("А.6", "Media upload flow"),
    ]:
        add_figure_placeholder(doc, number, title)

    add_section_heading(doc, "ДОДАТОК Б – Лістинги основних програмних модулів")
    add_plain_paragraph(
        doc,
        "У додатку Б можна розмістити повні або розширені версії файлів `backend/src/main.ts`, `backend/src/app.module.ts`, `backend/prisma/schema.prisma`, `frontend/contexts/AuthContext.tsx`, `frontend/components/CartModal.tsx`, `backend/src/orders/orders.service.ts`, `backend/src/media/media-storage.ts` і `backend/src/site-settings/site-settings.service.ts`.",
    )
    add_listing(
        doc,
        "Б.1",
        "Місце для повного лістингу Prisma schema",
        """
[МІСЦЕ ДЛЯ ПОВНОГО ЛІСТИНГУ backend/prisma/schema.prisma]
""",
    )
    add_listing(
        doc,
        "Б.2",
        "Місце для повного лістингу AuthContext",
        """
[МІСЦЕ ДЛЯ ПОВНОГО ЛІСТИНГУ frontend/contexts/AuthContext.tsx]
""",
    )

    add_section_heading(doc, "ДОДАТОК В – Додаткові скріншоти інтерфейсу")
    add_plain_paragraph(
        doc,
        "У додатку В доцільно розмістити додаткові стани інтерфейсу: empty state замовлень, validation errors, login challenge, verification required, auth unavailable, admin confirm dialogs і мобільні варіанти сторінок.",
    )
    for number, title in [
        ("В.1", "Стан порожньої історії замовлень"),
        ("В.2", "Login challenge у модальному вікні"),
        ("В.3", "Verification required після реєстрації"),
        ("В.4", "Auth unavailable page"),
        ("В.5", "Підтвердження видалення рослини"),
    ]:
        add_figure_placeholder(doc, number, title)

    add_section_heading(doc, "ДОДАТОК Г – Матеріали тестування та запуску системи")
    add_plain_paragraph(
        doc,
        "У додатку Г варто розмістити команди запуску frontend і backend, приклади `.env.example`, скріншоти результатів `npm test`, `npm run build`, а також умови запуску DB-backed smoke suite з `RUN_REAL_STACK_SMOKE=1`.",
    )
    add_table(
        doc,
        "Таблиця Г.1 – Команди для ручної перевірки PlantShop",
        ["Команда", "Призначення", "Примітка"],
        [
            ["cd frontend && npm run dev", "Запуск frontend на localhost:3000", "Потребує env"],
            ["cd backend && npm run dev", "Запуск backend на localhost:3001", "Потребує MySQL/env"],
            ["cd frontend && npm test", "Frontend tests", "Vitest"],
            ["cd backend && npm test -- --runInBand", "Backend tests", "Jest"],
            ["cd backend && npm run test:e2e", "E2E smoke suite", "Real-stack частина opt-in"],
        ],
        widths=[5.5, 6.2, 4.3],
    )


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_assignment(doc)
    add_annotation(doc)
    add_abstract(doc)
    add_toc(doc)
    add_intro(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_conclusions(doc)
    add_references(doc)
    add_appendices(doc)
    return doc


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT_FILE)
    print(f"DOCX generated: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
